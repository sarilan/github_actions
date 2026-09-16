#!/usr/bin/env python3
"""Construit le modèle Word du rapport mensuel (template-rapport-mensuel.docx).

Usage :
    python3 outils/rapport/build_template.py

Le modèle suit la structure de l'annexe C du document projet et contient des balises
docxtpl (Jinja2) remplies par generate_report.py :
résumé en 3 lignes, démarches réalisées ce mois, démarches en cours et prochaine étape,
échéances des 60 prochains jours, montants obtenus ou économisés, points d'attention pour
l'enfant, contacts et échanges avec le parent.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "template-rapport-mensuel.docx"
BLEU = RGBColor(0x1F, 0x3A, 0x5F)


def _field(paragraph, instr: str) -> None:
    run = paragraph.add_run()
    for tag, attrs, text in (("w:fldChar", {"w:fldCharType": "begin"}, None), ("w:instrText", {"xml:space": "preserve"}, instr),
                             ("w:fldChar", {"w:fldCharType": "separate"}, None), ("w:t", {}, "1"), ("w:fldChar", {"w:fldCharType": "end"}, None)):
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        if text:
            el.text = text
        run._r.append(el)


def _shade(cell, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for name, size, before in (("Title", 22, 0), ("Heading 1", 14, 14), ("Heading 2", 12, 10)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = BLEU
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True


def _page(doc: Document) -> None:
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0)
    s.top_margin = s.bottom_margin = Cm(1.8)
    hp = s.header.paragraphs[0]
    r = hp.add_run("RELAIS")
    r.bold = True
    r.font.color.rgb = BLEU
    r.font.size = Pt(11)
    hp.add_run("   ·   Rapport mensuel — {{ client.dossier }} — {{ periode.libelle }}").font.size = Pt(9)
    fp = s.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("Document confidentiel, réservé au client et au bénéficiaire.   Page ").font.size = Pt(8)
    _field(fp, "PAGE")
    fp.add_run(" / ").font.size = Pt(8)
    _field(fp, "NUMPAGES")
    for run in fp.runs:
        run.font.size = Pt(8)


def _row(table, texts: list[str], widths: list[float], bold: bool = False, shade: str | None = None, size: float = 9.5) -> None:
    row = table.add_row()
    for j, txt in enumerate(texts):
        c = row.cells[j]
        c.width = Cm(widths[j])
        c.text = ""
        run = c.paragraphs[0].add_run(txt)
        run.bold = bold
        run.font.size = Pt(size)
        if shade:
            _shade(c, shade)


def _table(doc: Document, headers: list[str], row_tag: str, cells: list[str], widths: list[float], empty_text: str) -> None:
    """Tableau docxtpl : en-tête, ligne « {%tr for %} », ligne de contenu, ligne « {%tr endfor %} »,
    puis une ligne affichée seulement si la liste est vide (« {%tr if not … %} »)."""
    table = doc.add_table(rows=0, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)
    for j, w in enumerate(widths):
        table.columns[j].width = Cm(w)
    n = len(headers)
    liste = row_tag.split(" in ")[1]
    _row(table, headers, widths, bold=True, shade="E8EEF5")
    _row(table, ["{%tr for " + row_tag + " %}"] + [""] * (n - 1), widths)
    _row(table, cells, widths)
    _row(table, ["{%tr endfor %}"] + [""] * (n - 1), widths)
    _row(table, ["{%tr if not " + liste + " %}"] + [""] * (n - 1), widths)
    _row(table, [empty_text] + [""] * (n - 1), widths)
    _row(table, ["{%tr endif %}"] + [""] * (n - 1), widths)
    doc.add_paragraph()


def build(path: Path = TEMPLATE) -> Path:
    doc = Document()
    _styles(doc)
    _page(doc)

    doc.add_paragraph("Rapport mensuel", style="Title")
    p = doc.add_paragraph()
    p.add_run("{{ periode.libelle }}").bold = True
    p.add_run(" — dossier {{ client.dossier }}")
    doc.add_paragraph("Bénéficiaire : {{ client.parent }}, {{ client.ville }}   ·   Destinataire : {{ client.enfant }}   ·   Interlocuteur Relais : {{ client.operateur }}")
    doc.add_paragraph("Édité le {{ periode.edite_le }}. Ce rapport reprend les démarches enregistrées dans l'espace partagé ; chaque action y est tracée avec ses pièces.")

    doc.add_heading("1. Résumé", level=1)
    doc.add_paragraph("{{ resume.ligne1 }}")
    doc.add_paragraph("{{ resume.ligne2 }}")
    doc.add_paragraph("{{ resume.ligne3 }}")

    doc.add_heading("2. Démarches réalisées ce mois", level=1)
    _table(doc, ["Organisme", "Action", "Résultat", "Terminée le"], "d in realisees",
           ["{{ d.organisme }}", "{{ d.titre }}", "{{ d.resultat }}", "{{ d.terminee_le }}"], [3.6, 4.6, 6.4, 2.4],
           "Aucune démarche clôturée ce mois.")

    doc.add_heading("3. Démarches en cours et prochaine étape", level=1)
    _table(doc, ["Organisme", "Démarche", "Statut", "Prochaine étape", "Échéance"], "d in en_cours",
           ["{{ d.organisme }}", "{{ d.titre }}", "{{ d.statut }}", "{{ d.prochaine_etape }}", "{{ d.echeance }}"], [3.2, 4.2, 2.8, 4.6, 2.2],
           "Aucune démarche en cours.")

    doc.add_heading("4. Échéances des 60 prochains jours", level=1)
    _table(doc, ["Date", "Organisme", "Objet", "Qui agit"], "e in echeances",
           ["{{ e.date }}", "{{ e.organisme }}", "{{ e.objet }}", "{{ e.qui }}"], [2.4, 3.8, 7.6, 3.2],
           "Aucune échéance dans les 60 prochains jours.")

    doc.add_heading("5. Montants obtenus ou économisés", level=1)
    _table(doc, ["Organisme", "Démarche", "Montant"], "m in montants",
           ["{{ m.organisme }}", "{{ m.titre }}", "{{ m.montant }}"], [4.5, 9.0, 3.5],
           "Aucun montant obtenu ce mois.")
    p = doc.add_paragraph()
    p.add_run("Total du mois : {{ montants_total }}   ·   Total depuis le début du mandat : {{ montants_cumul }}").bold = True

    doc.add_heading("6. Points d'attention pour vous", level=1)
    doc.add_paragraph("{%p for a in attention %}")
    doc.add_paragraph("{{ a.texte }}", style="List Bullet")
    doc.add_paragraph("{%p endfor %}")
    doc.add_paragraph("{%p if not attention %}")
    doc.add_paragraph("Aucune décision ni document ne sont attendus de votre part ce mois-ci.")
    doc.add_paragraph("{%p endif %}")

    doc.add_heading("7. Contacts et échanges avec votre parent", level=1)
    doc.add_paragraph("{{ contacts.texte }}")
    _table(doc, ["Reçu le", "Organisme", "Document", "Traitement"], "c in courriers",
           ["{{ c.recu_le }}", "{{ c.organisme }}", "{{ c.titre }}", "{{ c.traitement }}"], [2.4, 3.6, 6.0, 5.0],
           "Aucun courrier reçu ce mois.")

    p = doc.add_paragraph()
    p.add_run("Prochain rapport le 1er du mois prochain. Pour toute question : {{ contacts.email_relais }} — {{ contacts.telephone_relais }}.").italic = True
    doc.save(path)
    return path


if __name__ == "__main__":
    out = build()
    print(f"✔ {out.name} généré")
