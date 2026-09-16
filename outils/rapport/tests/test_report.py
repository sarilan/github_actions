"""Génération du rapport mensuel à partir des données de seed.py (mock de l'API Airtable)."""
import sys
from datetime import date
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
sys.path.insert(0, str(HERE.parent.parent / "airtable"))

import build_template  # noqa: E402
import generate_report as gr  # noqa: E402
from airtable_client import AirtableClient, AirtableError  # noqa: E402
from mock_api import MockAirtable  # noqa: E402
import create_base  # noqa: E402
import seed  # noqa: E402


@pytest.fixture(scope="module")
def base():
    mock = MockAirtable()
    api = AirtableClient(transport=mock.transport)
    base_id, _ = create_base.run(api, create_base.load_schema(), mock.workspace_id, None)
    seed.seed(api, base_id, log=lambda *_: None)
    return api, base_id


@pytest.fixture(scope="module")
def template(tmp_path_factory):
    return build_template.build(tmp_path_factory.mktemp("tpl") / "template.docx")


MOIS = date.today().strftime("%Y-%m")


def docx_text(path: Path) -> str:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    parts.append(doc.sections[0].header.paragraphs[0].text)
    return "\n".join(parts)


def test_template_is_built_and_committed():
    assert gr.TEMPLATE.exists(), "lancer build_template.py"
    text = docx_text(gr.TEMPLATE)
    for tag in ("{{ client.dossier }}", "{%tr for d in realisees %}", "{%tr for d in en_cours %}", "{%tr for e in echeances %}", "{%tr for m in montants %}", "{%p for a in attention %}", "{{ contacts.texte }}"):
        assert tag in text
    assert "RELAIS" in Document(gr.TEMPLATE).sections[0].header.paragraphs[0].text


def test_context_durand(base):
    api, base_id = base
    data = gr.fetch_data(api, base_id, "Durand — Sarcelles")
    ctx = gr.build_context(data, MOIS, edite_le=date.today())
    assert ctx["client"]["parent"] == "Monique Durand" and ctx["client"]["enfant"] == "Sarah Durand"
    titres_realisees = {d["titre"] for d in ctx["realisees"]}
    assert "Régularisation remboursements Ameli bloqués" in titres_realisees
    assert "Récupération accès FranceConnect et Ameli" not in titres_realisees  # clôturée le mois précédent
    assert {d["titre"] for d in ctx["en_cours"]} == {"Demande APA à domicile", "Comparatif mutuelle avant échéance", "Contestation facture EDF estimée"}
    assert ctx["montants_total"] == "952,40 €"
    assert ctx["montants_cumul"] == "952,40 €"
    assert any("Décision attendue de votre part" in a["texte"] and "mutuelle" in a["texte"].lower() for a in ctx["attention"])
    assert any(e["organisme"] == "Mutuelle Harmonie" for e in ctx["echeances"])  # date clé du contrat sous 60 jours
    assert all(e["qui"] in ("Vous", "Relais", "Votre parent") for e in ctx["echeances"])
    debut, fin = gr.periode_from(MOIS)
    attendus = [c for c in data["courriers"] if debut <= date.fromisoformat(c["Reçu le"]) <= fin]
    assert len(ctx["courriers"]) == len(attendus) >= 1
    assert "démarche" in ctx["resume"]["ligne1"] and "952,40" in ctx["resume"]["ligne2"]
    assert ctx["periode"]["libelle"].split(" ")[0].lower() == gr.MOIS_FR[date.today().month - 1]


def test_context_martin_critique(base):
    api, base_id = base
    ctx = gr.build_context(gr.fetch_data(api, base_id, "Martin — Lyon"), MOIS)
    assert any("Situation critique" in a["texte"] for a in ctx["attention"])
    assert any(e["qui"] == "Votre parent" for e in ctx["echeances"])  # renouvellement CNI en attente parent
    assert any("Action de votre parent" in a["texte"] for a in ctx["attention"])
    assert ctx["montants"] == [] and ctx["montants_total"] == "0,00 €"
    assert "Aucun montant obtenu" in ctx["resume"]["ligne2"]


def test_empty_month(base):
    api, base_id = base
    ctx = gr.build_context(gr.fetch_data(api, base_id, "Durand — Sarcelles"), "2020-01")
    assert ctx["realisees"] == [] and ctx["courriers"] == []
    assert ctx["montants_total"] == "0,00 €"
    assert ctx["montants_cumul"] == "952,40 €"


def test_generate_docx_and_pdf(base, template, tmp_path):
    api, base_id = base
    result = gr.generate(api, base_id, "Durand — Sarcelles", MOIS, tmp_path, edite_le=date.today(), template=template)
    assert result["docx"].exists() and result["pdf"].exists()
    assert result["docx"].name == f"rapport-durand-sarcelles-{MOIS}.docx"
    text = docx_text(result["docx"])
    import re
    assert not re.search(r"\{\{ |\{%", text), "balises docxtpl non rendues"
    for attendu in ("Monique Durand", "Sarah Durand", "Régularisation remboursements Ameli bloqués", "312,40", "640,00", "952,40", "Demande APA à domicile", "Points d'attention", "Contacts et échanges"):
        assert attendu in text, attendu
    assert "Aucune démarche clôturée ce mois." not in text
    reader = PdfReader(str(result["pdf"]))
    assert len(reader.pages) >= 1
    pdf_text = "\n".join(p.extract_text() or "" for p in reader.pages)
    assert "Durand" in pdf_text and "APA" in pdf_text
    assert result["pdf_engine"] in ("libreoffice", "reportlab")


def test_generate_empty_month_shows_placeholders(base, template, tmp_path):
    api, base_id = base
    result = gr.generate(api, base_id, "Martin — Lyon", "2020-01", tmp_path, template=template)
    text = docx_text(result["docx"])
    assert "Aucune démarche clôturée ce mois." in text
    assert "Aucun montant obtenu ce mois." in text
    assert "Aucun courrier reçu ce mois." in text


def test_pdf_fallback_reportlab(base, template, tmp_path, monkeypatch):
    api, base_id = base
    monkeypatch.setattr(gr, "_soffice", lambda: None)
    result = gr.generate(api, base_id, "Martin — Lyon", MOIS, tmp_path, template=template)
    assert result["pdf_engine"] == "reportlab"
    reader = PdfReader(str(result["pdf"]))
    assert "Robert Martin" in reader.pages[0].extract_text()


def test_pdf_requires_context_without_libreoffice(tmp_path, monkeypatch):
    monkeypatch.setattr(gr, "_soffice", lambda: None)
    with pytest.raises(RuntimeError):
        gr.docx_to_pdf(tmp_path / "x.docx", tmp_path / "x.pdf", None)


def test_unknown_client(base):
    api, base_id = base
    with pytest.raises(AirtableError):
        gr.fetch_data(api, base_id, "Inconnu — Nulle part")


def test_formatters():
    assert gr.fmt_montant(1234.5) == "1 234,50 €"
    assert gr.fmt_montant(-94) == "-94,00 €"
    assert gr.fmt_montant(None) == "—"
    assert gr.fmt_date("2026-09-03") == "3 septembre 2026"
    assert gr.fmt_date(None) == "—"
    assert gr.periode_from("2026-12") == (date(2026, 12, 1), date(2026, 12, 31))
    assert gr.periode_from("2028-02") == (date(2028, 2, 1), date(2028, 2, 29))
    assert gr.slug("Durand — Sarcelles") == "durand-sarcelles"


def test_cli_arguments(capsys, monkeypatch):
    assert gr.main(["--client", "X", "--mois", "2026-13"]) == 1
    monkeypatch.delenv("AIRTABLE_BASE_ID", raising=False)
    import airtable_client
    monkeypatch.setattr(airtable_client, "BASE_ID_FILE", Path("/nulle/part/.base_id"))
    assert gr.main(["--client", "X", "--mois", "2026-09"]) == 1
    assert "base inconnue" in capsys.readouterr().err
