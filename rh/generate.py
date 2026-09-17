#!/usr/bin/env python3
"""Génère rh/fiche-poste-operateur.docx depuis rh/sources/fiche-poste-operateur.md.

Usage :
    python3 rh/generate.py

Réutilise le convertisseur Markdown → docx de juridique/generate.py (mêmes styles,
en-tête « RELAIS », pagination, encadré de validation) et vérifie le document après
génération : 8 sections numérotées, 2 annexes, 3 courriers du test pratique présents.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "juridique"))

from generate import build_document, verify_document  # noqa: E402

SOURCE = HERE / "sources" / "fiche-poste-operateur.md"
OUTPUT = HERE / "fiche-poste-operateur.docx"
ARTICLES_ATTENDUS = 8


def verify_extra(path: Path) -> list[str]:
    doc = Document(path)
    problems: list[str] = []
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    annexes = [h for h in headings if h.startswith("Annexe")]
    if len(annexes) != 2:
        problems.append(f"{len(annexes)} annexes trouvées, 2 attendues")
    courriers = [p.text for p in doc.paragraphs if p.style.name == "Heading 3" and p.text.startswith("Courrier ")]
    if len(courriers) != 3:
        problems.append(f"{len(courriers)} courriers du test pratique, 3 attendus")
    texte = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    for attendu in ("900 à 1 200 €", "40 dossiers", "100 % à distance", "double contrôle", "Seuil de réussite"):
        if attendu not in texte:
            problems.append(f"mention absente : {attendu}")
    return problems


def main() -> int:
    n = build_document(SOURCE, OUTPUT)
    print(f"✔ {OUTPUT.name} généré ({n} sections)")
    problems = verify_document(OUTPUT, ARTICLES_ATTENDUS) + verify_extra(OUTPUT)
    for pb in problems:
        print(f"✘ {pb}")
    if problems:
        return 1
    print(f"✔ {OUTPUT.name} vérifié (8 sections, 2 annexes, 3 courriers, encadré, en-tête, pagination)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
