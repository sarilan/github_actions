#!/usr/bin/env python3
"""Génère les trois documents juridiques Word (.docx) depuis les sources Markdown.

Usage :
    python3 juridique/generate.py            # génère et vérifie les 3 documents
    python3 juridique/generate.py --check    # vérifie seulement les docx existants

Conventions des sources Markdown (juridique/sources/*.md) :
    # Titre                    titre du document (une seule fois, en tête)
    > texte                    encadré ; le premier encadré du document est l'avertissement
                               « à faire valider par un avocat », les suivants sont des modèles
                               de lettre encadrés
    ## Article — Libellé       article numéroté automatiquement (« Article 1 — Libellé »)
    ## Autre titre             titre de section non numéroté (Parties, Annexe…)
    ### Sous-titre             sous-titre
    - élément                  liste à puces (gras inline **…** accepté)
    | a | b |                  tableau (la première ligne est l'en-tête)
    [[SIGNATURES]]             bloc de signatures du mandat
    [[CHECKLIST]]              liste à cocher des organismes (annexe du mandat)
    paragraphe                 texte courant (gras inline **…** accepté)
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SOURCES = HERE / "sources"

# Nom du fichier de sortie -> (source Markdown, nombre d'articles numérotés attendu)
DOCUMENTS = {
    "mandat-administratif.docx": ("mandat-administratif.md", 14),
    "cgv.docx": ("cgv.md", 18),
    "politique-confidentialite.docx": ("politique-confidentialite.md", 12),
}

AVERTISSEMENT = "Document à faire valider par un avocat avant utilisation."

ORGANISMES_CHECKLIST = [
    "FranceConnect (identité numérique)",
    "Ameli — Assurance Maladie",
    "L'Assurance retraite (CNAV) / info-retraite.fr",
    "Agirc-Arrco (retraite complémentaire)",
    "MSA (régime agricole) ou autre régime de retraite",
    "impots.gouv.fr — Direction générale des Finances publiques",
    "ANTS — titres d'identité, certificat d'immatriculation, permis",
    "CAF — Caisse d'allocations familiales",
    "Conseil départemental — APA, aides à l'autonomie",
    "Mairie / CCAS",
    "Complémentaire santé (mutuelle) : ……………………………",
    "Assurance habitation : ……………………………",
    "Assurance automobile : ……………………………",
    "Fournisseur d'électricité : ……………………………",
    "Fournisseur de gaz : ……………………………",
    "Service des eaux : ……………………………",
    "Opérateur téléphonique : ……………………………",
    "Fournisseur d'accès internet : ……………………………",
    "Syndic de copropriété : ……………………………",
    "Banque (consultation uniquement) : ……………………………",
    "Seconde banque (consultation uniquement) : ……………………………",
    "La Poste — réexpédition du courrier",
    "Autre : ……………………………",
    "Autre : ……………………………",
]

_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


# ---------------------------------------------------------------------------
# Mise en forme
# ---------------------------------------------------------------------------
def _set_cell_shading(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _set_cell_borders(cell, color: str = "1F3A5F", size: int = 12) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tc_pr.append(borders)


def _fix_layout(table) -> None:
    """Impose une largeur de colonnes fixe (Word ignore sinon les largeurs des cellules)."""
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)


def _keep_table_together(table) -> None:
    """Empêche le tableau de se couper entre deux pages."""
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    rows = table.rows
    for idx, row in enumerate(rows):
        for cell in row.cells:
            for par in cell.paragraphs:
                par.paragraph_format.keep_together = True
                if idx < len(rows) - 1:
                    par.paragraph_format.keep_with_next = True


def _add_field(paragraph, instr: str) -> None:
    """Insère un champ Word (PAGE, NUMPAGES…) dans un paragraphe."""
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for el in (fld_begin, instr_el, fld_sep, txt, fld_end):
        run._r.append(el)


def _configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    # Police est-asiatique alignée pour éviter les substitutions.
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

    for name, size, before in (("Heading 1", 16, 18), ("Heading 2", 13, 14), ("Heading 3", 11.5, 10)):
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True


def _configure_page(doc: Document, header_text: str) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.text = ""
    run = hp.add_run("RELAIS")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    hp.add_run("   ·   " + header_text).font.size = Pt(9)
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("Page ").font.size = Pt(9)
    _add_field(fp, "PAGE")
    fp.add_run(" / ").font.size = Pt(9)
    _add_field(fp, "NUMPAGES")
    for r in fp.runs:
        r.font.size = Pt(9)


def _add_inline(paragraph, text: str, bold_all: bool = False) -> None:
    """Ajoute du texte avec prise en charge du gras inline **…**."""
    pos = 0
    for m in _BOLD_RE.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos:m.start()])
            r.bold = bold_all
        r = paragraph.add_run(m.group(1))
        r.bold = True
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        r.bold = bold_all


def _add_boxed(doc: Document, text: str, warning: bool) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    _set_cell_shading(cell, "FFF4D6" if warning else "F3F6FA")
    _set_cell_borders(cell, color="C77700" if warning else "1F3A5F", size=16 if warning else 8)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if warning else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    _add_inline(p, text, bold_all=warning)
    if warning:
        for r in p.runs:
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(0x8A, 0x4B, 0x00)
    doc.add_paragraph()


def _add_table(doc: Document, rows: list[list[str]]) -> None:
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            txt = row[j] if j < len(row) else ""
            _add_inline(p, txt, bold_all=(i == 0))
            for r in p.runs:
                r.font.size = Pt(9.5)
            if i == 0:
                _set_cell_shading(cell, "E8EEF5")
    doc.add_paragraph()


def _add_signatures(doc: Document) -> None:
    p = doc.add_paragraph()
    p.add_run("Fait à ……………………………………, le ……… / ……… / ………………, en deux exemplaires originaux.")
    p.paragraph_format.space_before = Pt(10)

    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = table.rows[0].cells
    left.paragraphs[0].add_run("Le Mandant").bold = True
    left.add_paragraph("Nom, prénom : ……………………………………")
    q = left.add_paragraph()
    q.add_run("Mention manuscrite obligatoire : « Bon pour mandat »").italic = True
    left.add_paragraph("\n\n\nSignature :")
    right.paragraphs[0].add_run("Le Mandataire").bold = True
    right.add_paragraph("{{MANDATAIRE_RAISON_SOCIALE}} — service Relais")
    right.add_paragraph("Représenté par : {{MANDATAIRE_REPRESENTANT}}")
    right.add_paragraph("\n\n\nSignature et cachet :")
    cp = table.rows[1].cells[0].merge(table.rows[1].cells[1])
    cp.paragraphs[0].add_run("Copie reçue par la Personne de confiance").bold = True
    cp.add_paragraph("Nom, prénom : ……………………………………   Date : ……… / ……… / ………………")
    cp.add_paragraph("Je reconnais avoir reçu copie du présent mandat et accepte d'être destinataire des rapports et des alertes.")
    cp.add_paragraph("\n\nSignature :")
    for row in table.rows:
        for cell in row.cells:
            for par in cell.paragraphs:
                par.paragraph_format.space_after = Pt(2)
    _keep_table_together(table)
    p.paragraph_format.keep_with_next = True
    doc.add_paragraph()


def _add_checklist(doc: Document) -> None:
    rows = [["", "Organisme ou compte", "Identifiant, n° de dossier ou remarque"]]
    for org in ORGANISMES_CHECKLIST:
        rows.append(["☐", org, ""])
    table = doc.add_table(rows=len(rows), cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _fix_layout(table)
    widths = (Cm(1.0), Cm(9.0), Cm(6.6))
    for j, w in enumerate(widths):
        table.columns[j].width = w
    for i, row in enumerate(rows):
        for j, txt in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.width = widths[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(txt)
            r.font.size = Pt(9.5)
            r.bold = i == 0
            if j == 0 and i > 0:
                r.font.size = Pt(13)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if i == 0:
                _set_cell_shading(cell, "E8EEF5")
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Paraphe du Mandant : ……………   Paraphe du Mandataire : ……………").italic = True


# ---------------------------------------------------------------------------
# Conversion Markdown -> docx
# ---------------------------------------------------------------------------
def _flush_paragraph(doc: Document, buf: list[str]) -> None:
    if buf:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _add_inline(p, " ".join(s.strip() for s in buf))
        buf.clear()


def build_document(source: Path, output: Path) -> int:
    """Construit le docx et retourne le nombre d'articles numérotés produits."""
    lines = source.read_text(encoding="utf-8").splitlines()
    doc = Document()
    _configure_styles(doc)

    title = next((ln[2:].strip() for ln in lines if ln.startswith("# ")), source.stem)
    _configure_page(doc, title)
    doc.add_paragraph(title, style="Title")

    article_no = 0
    warning_done = False
    para_buf: list[str] = []
    table_buf: list[list[str]] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("|"):
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                table_buf.append(cells)
            i += 1
            continue
        if table_buf:
            _flush_paragraph(doc, para_buf)
            _add_table(doc, table_buf)
            table_buf = []

        if not stripped:
            _flush_paragraph(doc, para_buf)
        elif stripped.startswith("# "):
            pass  # titre déjà posé
        elif stripped.startswith("## "):
            _flush_paragraph(doc, para_buf)
            heading = stripped[3:].strip()
            if heading.startswith("Article"):
                article_no += 1
                label = heading.split("—", 1)[1].strip() if "—" in heading else heading
                doc.add_heading(f"Article {article_no} — {label}", level=2)
            elif heading.startswith("Annexe"):
                doc.add_page_break()
                doc.add_heading(heading, level=1)
            else:
                doc.add_heading(heading, level=1)
        elif stripped.startswith("### "):
            _flush_paragraph(doc, para_buf)
            doc.add_heading(stripped[4:].strip(), level=3)
        elif stripped.startswith("> "):
            _flush_paragraph(doc, para_buf)
            text = stripped[2:].strip()
            is_warning = not warning_done and text == AVERTISSEMENT
            _add_boxed(doc, text, warning=is_warning)
            warning_done = warning_done or is_warning
        elif stripped.startswith("- "):
            _flush_paragraph(doc, para_buf)
            p = doc.add_paragraph(style="List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _add_inline(p, stripped[2:].strip())
        elif stripped == "[[SIGNATURES]]":
            _flush_paragraph(doc, para_buf)
            _add_signatures(doc)
        elif stripped == "[[CHECKLIST]]":
            _flush_paragraph(doc, para_buf)
            _add_checklist(doc)
        else:
            para_buf.append(stripped)
        i += 1

    _flush_paragraph(doc, para_buf)
    if table_buf:
        _add_table(doc, table_buf)

    if not warning_done:
        raise RuntimeError(f"{source.name} : l'encadré « {AVERTISSEMENT} » est absent")

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return article_no


# ---------------------------------------------------------------------------
# Vérification
# ---------------------------------------------------------------------------
def verify_document(path: Path, expected_articles: int) -> list[str]:
    """Réouvre le docx et retourne la liste des anomalies (vide si tout va bien)."""
    problems: list[str] = []
    doc = Document(path)
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 2" and p.text.startswith("Article ")]
    if len(headings) != expected_articles:
        problems.append(f"{path.name} : {len(headings)} articles trouvés, {expected_articles} attendus")
    numbers = [int(h.split(" ")[1]) for h in headings]
    if numbers != list(range(1, len(numbers) + 1)):
        problems.append(f"{path.name} : numérotation des articles discontinue : {numbers}")

    first_page_text = " ".join(
        cell.text for table in doc.tables[:1] for row in table.rows for cell in row.cells
    )
    if AVERTISSEMENT not in first_page_text:
        problems.append(f"{path.name} : encadré d'avertissement absent de la première page")

    header_text = doc.sections[0].header.paragraphs[0].text
    if "RELAIS" not in header_text:
        problems.append(f"{path.name} : en-tête « RELAIS » absent")
    footer_xml = doc.sections[0].footer._element.xml
    if "PAGE" not in footer_xml or "NUMPAGES" not in footer_xml:
        problems.append(f"{path.name} : pagination absente du pied de page")

    full_text = "\n".join(p.text for p in doc.paragraphs) + "\n" + " ".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    for token in ("TODO", "à compléter", "XXX", "lorem"):
        if token.lower() in full_text.lower():
            problems.append(f"{path.name} : marqueur interdit « {token} » présent")
    return problems


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--check", action="store_true", help="vérifie seulement les documents existants")
    parser.add_argument("--out", type=Path, default=HERE, help="dossier de sortie (défaut : juridique/)")
    args = parser.parse_args(argv)

    all_problems: list[str] = []
    for output_name, (source_name, expected) in DOCUMENTS.items():
        output = args.out / output_name
        if not args.check:
            produced = build_document(SOURCES / source_name, output)
            print(f"✔ {output_name} généré ({produced} articles)")
        problems = verify_document(output, expected)
        if problems:
            all_problems.extend(problems)
            for pb in problems:
                print(f"✘ {pb}")
        else:
            print(f"✔ {output_name} vérifié ({expected} articles, encadré, en-tête, pagination)")

    if all_problems:
        print(f"\n{len(all_problems)} anomalie(s).")
        return 1
    print("\nTous les documents sont conformes.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
