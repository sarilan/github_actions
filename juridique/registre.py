#!/usr/bin/env python3
"""Génère le registre des activités de traitement (article 30 RGPD) au format Excel.

Usage :
    python3 juridique/registre.py

Les colonnes reprennent le modèle de registre simplifié publié par la CNIL :
une feuille « Responsable » (identité du responsable de traitement), une feuille
« Registre » (une ligne par traitement) et une feuille « Sous-traitants ».
"""
from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
OUTPUT = HERE / "registre-traitements-rgpd.xlsx"

COLONNES = [
    "N°",
    "Nom du traitement",
    "Finalité(s) principale(s)",
    "Sous-finalités",
    "Base légale",
    "Catégories de personnes concernées",
    "Catégories de données",
    "Données sensibles (art. 9) ou NIR",
    "Source des données",
    "Destinataires internes",
    "Destinataires externes",
    "Sous-traitants",
    "Transferts hors UE",
    "Durée de conservation (base active)",
    "Durée d'archivage intermédiaire",
    "Mesures de sécurité",
    "Service responsable",
    "Analyse d'impact (AIPD)",
    "Date de création",
    "Date de dernière mise à jour",
]

TRAITEMENTS = [
    {
        "Nom du traitement": "Gestion des mandats et des dossiers clients",
        "Finalité(s) principale(s)": "Conclure, exécuter et suivre le mandat de représentation administrative et l'abonnement au service Relais.",
        "Sous-finalités": "Vérification d'identité ; diagnostic initial ; suivi des démarches ; rapports mensuels et alertes ; relation avec la personne de confiance.",
        "Base légale": "Exécution du contrat et du mandat (art. 6.1.b RGPD) ; consentement explicite du bénéficiaire pour les données de santé (art. 9.2.a RGPD).",
        "Catégories de personnes concernées": "Bénéficiaires (parents) ; clients (enfants) ; personnes de confiance ; tiers mentionnés dans les dossiers.",
        "Catégories de données": "Identité, pièces d'identité, coordonnées, lien de parenté, situation familiale, administrative, sociale et fiscale, contrats, décisions des organismes, échéances, contenu des rapports.",
        "Données sensibles (art. 9) ou NIR": "Oui : données de santé strictement nécessaires aux démarches (Assurance Maladie, mutuelle, APA, hospitalisation, EHPAD) ; NIR dans les cas prévus.",
        "Source des données": "Bénéficiaire, client, organismes tiers (courriers, espaces en ligne).",
        "Destinataires internes": "Responsable du service ; opérateur en charge du dossier.",
        "Destinataires externes": "Organismes destinataires des démarches ; personne de confiance ; assureur RC pro et médiateur en cas de litige.",
        "Sous-traitants": "Airtable (base de gestion) ; service de messagerie ; service de signature électronique.",
        "Transferts hors UE": "Envoi des rapports au client dans son pays de résidence (art. 49.1.b) ; sous-traitants hors UE encadrés par clauses contractuelles types le cas échéant.",
        "Durée de conservation (base active)": "Durée du mandat.",
        "Durée d'archivage intermédiaire": "3 ans après la fin du mandat (preuve), puis suppression ; pièces médicales supprimées dès la clôture de la démarche.",
        "Mesures de sécurité": "Accès nominatifs ; authentification à deux facteurs ; chiffrement TLS ; journal d'accès ; engagement de confidentialité des opérateurs ; sauvegardes.",
        "Service responsable": "Direction du service Relais.",
        "Analyse d'impact (AIPD)": "Oui, requise (données de santé de personnes vulnérables) — à réaliser avant le premier client.",
    },
    {
        "Nom du traitement": "Réception, numérisation et traitement du courrier",
        "Finalité(s) principale(s)": "Recevoir, numériser, classer et traiter le courrier administratif du bénéficiaire dans le délai de 48 h ouvrées.",
        "Sous-finalités": "Réexpédition postale ; classification automatique (organisme, type, urgence) ; création des démarches ; renvoi du courrier personnel au bénéficiaire.",
        "Base légale": "Exécution du contrat et du mandat (art. 6.1.b RGPD) ; consentement explicite pour les données de santé contenues dans le courrier (art. 9.2.a RGPD).",
        "Catégories de personnes concernées": "Bénéficiaires ; expéditeurs des courriers ; tiers mentionnés.",
        "Catégories de données": "Contenu des courriers : identité, références de dossier, montants, dates limites, décisions, relevés, factures.",
        "Données sensibles (art. 9) ou NIR": "Oui, possibles : courriers de l'Assurance Maladie, de la mutuelle, des établissements de santé ; NIR.",
        "Source des données": "Courrier postal du bénéficiaire ; photographies transmises par le bénéficiaire ou un proche.",
        "Destinataires internes": "Opérateur en charge du dossier ; responsable du service.",
        "Destinataires externes": "Bénéficiaire (renvoi du courrier personnel) ; client (rapport).",
        "Sous-traitants": "Service de domiciliation et de numérisation du courrier ; Airtable ; hébergeur des fichiers numérisés.",
        "Transferts hors UE": "Aucun par principe ; hébergement UE exigé du service de numérisation.",
        "Durée de conservation (base active)": "Durée du mandat.",
        "Durée d'archivage intermédiaire": "3 ans après la fin du mandat, puis suppression ; originaux papier renvoyés au bénéficiaire ou détruits sous 3 mois.",
        "Mesures de sécurité": "Chiffrement des fichiers au repos et en transit ; accès nominatifs ; destruction sécurisée des originaux ; contrat art. 28 avec le service de numérisation.",
        "Service responsable": "Opérations.",
        "Analyse d'impact (AIPD)": "Couverte par l'AIPD du traitement n° 1.",
    },
    {
        "Nom du traitement": "Gestion des accès aux comptes en ligne et journal d'accès",
        "Finalité(s) principale(s)": "Conserver de manière sécurisée les identifiants des comptes en ligne du bénéficiaire et tracer chaque accès effectué par un opérateur.",
        "Sous-finalités": "Récupération ou création des comptes ; réception des codes de vérification sur la ligne dédiée ; révocation des accès en fin de mandat ; détection d'accès anormal.",
        "Base légale": "Exécution du mandat (art. 6.1.b RGPD) ; intérêt légitime à la sécurité des accès pour le journal (art. 6.1.f RGPD).",
        "Catégories de personnes concernées": "Bénéficiaires ; opérateurs (journal d'accès).",
        "Catégories de données": "Identifiants et mots de passe ; codes de vérification ; adresse et téléphone déclarés aux organismes ; horodatage des accès et identité de l'opérateur.",
        "Données sensibles (art. 9) ou NIR": "Non directement ; l'accès aux comptes donne accès à des données de santé (Ameli) traitées au titre du traitement n° 1.",
        "Source des données": "Bénéficiaire ; organismes (réinitialisation) ; système.",
        "Destinataires internes": "Opérateur habilité sur le dossier ; responsable du service (journal).",
        "Destinataires externes": "Aucun.",
        "Sous-traitants": "Bitwarden (coffre-fort chiffré) ; opérateur de téléphonie (ligne dédiée).",
        "Transferts hors UE": "Selon la localisation du coffre-fort : hébergement UE privilégié, sinon clauses contractuelles types.",
        "Durée de conservation (base active)": "Identifiants : durée du mandat, suppression sous 15 jours après la fin. Journal d'accès : 1 an.",
        "Durée d'archivage intermédiaire": "Aucun pour les identifiants ; journal conservé 1 an puis supprimé.",
        "Mesures de sécurité": "Chiffrement de bout en bout ; un coffre par client ; accès nominatifs avec double authentification ; interdiction de transmettre un identifiant par courriel ; révocation immédiate au départ d'un opérateur.",
        "Service responsable": "Opérations / sécurité.",
        "Analyse d'impact (AIPD)": "Oui, incluse dans l'AIPD globale.",
    },
    {
        "Nom du traitement": "Facturation, encaissement et comptabilité",
        "Finalité(s) principale(s)": "Facturer les abonnements et forfaits, encaisser les paiements, tenir la comptabilité et répondre aux obligations fiscales.",
        "Sous-finalités": "Détermination du régime de TVA selon le pays du client ; relances d'impayés ; gestion des rétractations et remboursements.",
        "Base légale": "Exécution du contrat (art. 6.1.b RGPD) ; obligation légale comptable et fiscale (art. 6.1.c RGPD).",
        "Catégories de personnes concernées": "Clients.",
        "Catégories de données": "Identité, adresse de facturation, pays de résidence, courriel, historique des paiements, jeton de paiement (jamais le numéro complet de carte), factures.",
        "Données sensibles (art. 9) ou NIR": "Non.",
        "Source des données": "Client ; prestataire de paiement.",
        "Destinataires internes": "Direction ; personne en charge de la facturation.",
        "Destinataires externes": "Expert-comptable ; administration fiscale ; prestataire de paiement.",
        "Sous-traitants": "Prestataire de paiement ; logiciel de facturation ; expert-comptable (responsable de traitement distinct pour sa mission).",
        "Transferts hors UE": "Selon le prestataire de paiement, encadrés par clauses contractuelles types.",
        "Durée de conservation (base active)": "Durée de la relation contractuelle.",
        "Durée d'archivage intermédiaire": "10 ans pour les factures et pièces comptables (art. L. 123-22 Code de commerce).",
        "Mesures de sécurité": "Paiement délégué à un prestataire certifié PCI-DSS ; accès restreint ; sauvegardes.",
        "Service responsable": "Administration / finances.",
        "Analyse d'impact (AIPD)": "Non requise.",
    },
    {
        "Nom du traitement": "Prospection et gestion des demandes de contact",
        "Finalité(s) principale(s)": "Répondre aux demandes reçues par le formulaire du site, organiser l'appel de présentation et informer les prospects et clients des évolutions du service.",
        "Sous-finalités": "Relances des prospects sans réponse (J+3, J+10) ; parrainage ; réunions d'information ; statistiques anonymisées de conversion.",
        "Base légale": "Mesures précontractuelles à la demande du prospect (art. 6.1.b RGPD) ; consentement pour l'envoi d'informations aux prospects ; intérêt légitime pour les clients existants (art. 6.1.f RGPD).",
        "Catégories de personnes concernées": "Prospects ; clients ; parrains et filleuls.",
        "Catégories de données": "Nom, courriel, téléphone, pays de résidence, ville et âge du parent, démarche posant problème, créneau d'appel, historique des échanges.",
        "Données sensibles (art. 9) ou NIR": "Non recherchées ; les informations spontanément fournies dans le formulaire (ex. hospitalisation) sont limitées au strict nécessaire et supprimées si aucune souscription.",
        "Source des données": "Prospect (formulaire, courriel, appel) ; parrain.",
        "Destinataires internes": "Direction ; personne en charge des appels.",
        "Destinataires externes": "Aucun.",
        "Sous-traitants": "Formspree (réception du formulaire) ; service de messagerie ; outil de visioconférence.",
        "Transferts hors UE": "Formspree : États-Unis, encadré par clauses contractuelles types ; envoi des courriels au prospect dans son pays de résidence.",
        "Durée de conservation (base active)": "Jusqu'à la souscription ou 3 ans après le dernier contact.",
        "Durée d'archivage intermédiaire": "Aucun ; suppression à l'issue de la durée.",
        "Mesures de sécurité": "Formulaire sans cookie ; accès restreint à la boîte de réception ; désinscription possible à chaque envoi ; aucun traceur publicitaire sur le site.",
        "Service responsable": "Commercial.",
        "Analyse d'impact (AIPD)": "Non requise.",
    },
]

SOUS_TRAITANTS = [
    ("Airtable", "Base de gestion des dossiers", "{{AIRTABLE_ENTITE}}", "{{AIRTABLE_HEBERGEMENT}}", "Contrat art. 28 (conditions du fournisseur) ; clauses contractuelles types si hors UE"),
    ("Bitwarden", "Coffre-fort d'identifiants", "{{BITWARDEN_ENTITE}}", "{{BITWARDEN_HEBERGEMENT}}", "Contrat art. 28 ; chiffrement de bout en bout"),
    ("Service de numérisation du courrier", "Domiciliation, scan, réexpédition", "{{SERVICE_SCAN_NOM}}", "{{SERVICE_SCAN_HEBERGEMENT}}", "Contrat art. 28 ; hébergement UE exigé"),
    ("Téléphonie", "Ligne dédiée, codes de vérification", "{{TELEPHONIE_NOM}}", "{{TELEPHONIE_HEBERGEMENT}}", "Contrat art. 28"),
    ("Messagerie", "Courriels, rapports", "{{MESSAGERIE_NOM}}", "{{MESSAGERIE_HEBERGEMENT}}", "Contrat art. 28"),
    ("Paiement", "Encaissement", "{{PAIEMENT_NOM}}", "{{PAIEMENT_HEBERGEMENT}}", "Contrat art. 28 ; PCI-DSS"),
    ("Formspree", "Formulaire de contact", "{{FORMSPREE_ENTITE}}", "{{FORMSPREE_HEBERGEMENT}}", "Clauses contractuelles types (États-Unis)"),
]

RESPONSABLE = [
    ("Responsable de traitement", "{{MANDATAIRE_RAISON_SOCIALE}}"),
    ("Forme juridique", "{{MANDATAIRE_FORME_JURIDIQUE}}"),
    ("SIREN", "{{MANDATAIRE_SIREN}}"),
    ("Adresse du siège", "{{MANDATAIRE_ADRESSE}}"),
    ("Représentant légal", "{{MANDATAIRE_REPRESENTANT}}"),
    ("Référent protection des données", "{{DPO_EMAIL}}"),
    ("Service exploité", "Relais — mandataire administratif à distance"),
    ("Date d'établissement du registre", date.today().isoformat()),
    ("Avertissement", "Document à faire valider par un avocat avant utilisation."),
]

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="B7C4D3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def _style_header(ws, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
    ws.freeze_panes = "C2" if ncols > 3 else "A2"


def build(output: Path = OUTPUT) -> Path:
    wb = Workbook()

    ws0 = wb.active
    ws0.title = "Responsable"
    ws0.append(["Élément", "Valeur"])
    for k, v in RESPONSABLE:
        ws0.append([k, v])
    _style_header(ws0, 2)
    ws0.column_dimensions["A"].width = 38
    ws0.column_dimensions["B"].width = 70
    for row in ws0.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
            c.border = BORDER

    ws = wb.create_sheet("Registre")
    ws.append(COLONNES)
    today = date.today().isoformat()
    for n, t in enumerate(TRAITEMENTS, start=1):
        row = [n]
        for col in COLONNES[1:]:
            if col == "Date de création" or col == "Date de dernière mise à jour":
                row.append(today)
            else:
                row.append(t[col])
        ws.append(row)
    _style_header(ws, len(COLONNES))
    widths = [5, 30, 40, 40, 40, 32, 45, 35, 30, 28, 32, 32, 35, 30, 35, 45, 22, 30, 14, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
            c.border = BORDER
        ws.row_dimensions[row[0].row].height = 190

    ws2 = wb.create_sheet("Sous-traitants")
    ws2.append(["Sous-traitant", "Rôle", "Entité juridique", "Hébergement des données", "Encadrement contractuel"])
    for r in SOUS_TRAITANTS:
        ws2.append(list(r))
    _style_header(ws2, 5)
    for col, w in zip("ABCDE", (32, 34, 30, 30, 50)):
        ws2.column_dimensions[col].width = w
    for row in ws2.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
            c.border = BORDER

    wb.save(output)
    return output


def verify(path: Path = OUTPUT) -> list[str]:
    problems: list[str] = []
    wb = load_workbook(path)
    if wb.sheetnames != ["Responsable", "Registre", "Sous-traitants"]:
        problems.append(f"feuilles inattendues : {wb.sheetnames}")
    ws = wb["Registre"]
    header = [c.value for c in ws[1]]
    if header != COLONNES:
        problems.append("en-têtes du registre différents du modèle")
    if ws.max_row - 1 != len(TRAITEMENTS):
        problems.append(f"{ws.max_row - 1} traitements, {len(TRAITEMENTS)} attendus")
    for row in ws.iter_rows(min_row=2, values_only=True):
        empties = [COLONNES[i] for i, v in enumerate(row) if v in (None, "")]
        if empties:
            problems.append(f"traitement n° {row[0]} : cellules vides {empties}")
    return problems


if __name__ == "__main__":
    out = build()
    pbs = verify(out)
    if pbs:
        for pb in pbs:
            print("✘", pb)
        sys.exit(1)
    print(f"✔ {out.name} généré et vérifié ({len(TRAITEMENTS)} traitements, {len(COLONNES)} colonnes)")
